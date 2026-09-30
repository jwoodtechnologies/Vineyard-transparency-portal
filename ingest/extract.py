"""Deterministic text extraction with page boundaries (Poppler, python-docx, openpyxl)."""
from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
from dataclasses import dataclass, field

# A PDF page with less extractable text than this is probably a scan.
MIN_CHARS_PER_PAGE = 40


@dataclass
class Page:
    number: int | None  # None for formats without pages (docx, xlsx, csv, txt)
    text: str
    section: str | None = None
    ocr: bool = False


@dataclass
class Extraction:
    pages: list[Page] = field(default_factory=list)
    page_count: int | None = None
    status: str = "extracted"  # extracted | empty | unsupported | failed
    ocr_status: str = "not_required"  # not_required | needed | complete | failed
    detail: str | None = None

    @property
    def characters(self) -> int:
        return sum(len(p.text) for p in self.pages)


def sniff(path: str) -> str | None:
    with open(path, "rb") as f:
        head = f.read(8)
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"
    return None


def _run(cmd: list[str], timeout: int = 300) -> str:
    return subprocess.run(cmd, check=True, capture_output=True, timeout=timeout).stdout.decode("utf-8", "replace")


def pdf_page_count(path: str) -> int | None:
    try:
        out = _run(["pdfinfo", path], timeout=60)
    except (subprocess.SubprocessError, OSError):
        return None
    m = re.search(r"^Pages:\s+(\d+)", out, re.M)
    return int(m.group(1)) if m else None


def extract_pdf(path: str, ocr: bool = False, ocr_max_pages: int = 30) -> Extraction:
    count = pdf_page_count(path)
    try:
        raw = _run(["pdftotext", "-enc", "UTF-8", "-eol", "unix", path, "-"], timeout=600)
    except subprocess.TimeoutExpired:
        return Extraction(status="failed", page_count=count, detail="pdftotext timed out")
    except (subprocess.SubprocessError, OSError) as e:
        return Extraction(status="failed", page_count=count, detail=f"pdftotext failed: {e}")
    texts = raw.split("\f")
    if texts and not texts[-1].strip():
        texts = texts[:-1]
    pages = [Page(number=i + 1, text=clean(t)) for i, t in enumerate(texts)]
    count = count or len(pages)
    low = [p for p in pages if len(p.text) < MIN_CHARS_PER_PAGE]
    ex = Extraction(pages=[p for p in pages if p.text], page_count=count)
    if count and len(low) / max(1, count) > 0.5:
        ex.ocr_status = "needed"
        if ocr and shutil.which("tesseract") and shutil.which("pdftoppm"):
            ex = _ocr_pages(path, ex, [p.number for p in low][:ocr_max_pages])
    if not ex.pages:
        ex.status = "empty"
    return ex


def _ocr_pages(path: str, ex: Extraction, numbers: list[int | None]) -> Extraction:
    import os
    import tempfile

    done = {p.number: p for p in ex.pages}
    ok = 0
    with tempfile.TemporaryDirectory() as tmp:
        for n in numbers:
            if n is None:
                continue
            try:
                subprocess.run(["pdftoppm", "-r", "200", "-gray", "-f", str(n), "-l", str(n), "-png", path, f"{tmp}/p"], check=True, capture_output=True, timeout=120)
                img = next((f"{tmp}/{f}" for f in os.listdir(tmp) if f.endswith(".png")), None)
                if not img:
                    continue
                text = _run(["tesseract", img, "-", "--psm", "3"], timeout=180)
                os.unlink(img)
            except (subprocess.SubprocessError, OSError):
                continue
            text = clean(text)
            if text:
                done[n] = Page(number=n, text=text, ocr=True)
                ok += 1
    ex.pages = [done[k] for k in sorted(k for k in done if k is not None)]
    ex.ocr_status = "complete" if ok else "failed"
    return ex


def extract_docx(path: str) -> Extraction:
    import docx  # python-docx

    try:
        d = docx.Document(path)
    except Exception as e:  # noqa: BLE001 - malformed files are data, not bugs
        return Extraction(status="failed", detail=f"docx: {e}")
    pages: list[Page] = []
    section = None
    buf: list[str] = []

    def flush():
        if buf:
            pages.append(Page(number=None, text=clean("\n".join(buf)), section=section))
            buf.clear()

    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.lower().startswith(("heading", "title")):
            flush()
            section = t[:200]
        buf.append(t)
    flush()
    for table in d.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        text = clean("\n".join(r for r in rows if r.strip(" |")))
        if text:
            pages.append(Page(number=None, text=text, section="Table"))
    return Extraction(pages=[p for p in pages if p.text], status="extracted" if pages else "empty")


def extract_xlsx(path: str, max_chars: int = 400_000) -> Extraction:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        return Extraction(status="failed", detail=f"xlsx: {e}")
    pages: list[Page] = []
    total = 0
    for ws in wb.worksheets:
        lines = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))
                total += len(lines[-1])
            if total > max_chars:
                break
        if lines:
            pages.append(Page(number=None, text=clean("\n".join(lines)), section=f"Sheet: {ws.title}"[:200]))
        if total > max_chars:
            break
    wb.close()
    return Extraction(pages=pages, status="extracted" if pages else "empty")


def extract_text_file(path: str, is_csv: bool) -> Extraction:
    data = open(path, "rb").read(5_000_000)
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if is_csv:
        rows = [" | ".join(c.strip() for c in row if c.strip()) for row in csv.reader(io.StringIO(text))]
        text = "\n".join(r for r in rows if r)
    text = clean(text)
    return Extraction(pages=[Page(number=None, text=text)] if text else [], status="extracted" if text else "empty")


def extract_legacy(path: str, kind: str) -> Extraction:
    tool = {"doc": ["antiword", path], "xls": ["xls2csv", path], "ppt": ["catppt", path]}.get(kind)
    if not tool or not shutil.which(tool[0]):
        return Extraction(status="unsupported", detail=f"No extractor installed for .{kind}")
    try:
        text = clean(_run(tool, timeout=120))
    except (subprocess.SubprocessError, OSError) as e:
        return Extraction(status="failed", detail=str(e))
    return Extraction(pages=[Page(number=None, text=text)] if text else [], status="extracted" if text else "empty")


def extract(path: str, ext: str, mime: str, ocr: bool = False) -> Extraction:
    kind = sniff(path)
    if kind == "pdf" or ext == "pdf" or mime == "application/pdf":
        return extract_pdf(path, ocr=ocr)
    if ext == "docx" or (kind == "zip" and "wordprocessing" in mime):
        return extract_docx(path)
    if ext == "xlsx" or (kind == "zip" and "spreadsheet" in mime):
        return extract_xlsx(path)
    if ext in ("csv",) or mime == "text/csv":
        return extract_text_file(path, is_csv=True)
    if ext in ("txt",) or mime == "text/plain":
        return extract_text_file(path, is_csv=False)
    if ext in ("doc", "xls", "ppt"):
        return extract_legacy(path, ext)
    return Extraction(status="unsupported", detail=f"Unsupported type .{ext or '?'} ({mime or 'unknown'})")


_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def clean(text: str) -> str:
    text = _CONTROL.sub(" ", text.replace("\r\n", "\n").replace("\r", "\n"))
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_html(path: str, kind: str | None = None) -> tuple[Extraction, str | None]:
    """Main readable text of a web page (city page or Sheriff press release) and its heading.
    Menus, headers, footers, scripts and sidebars are dropped so only the page's own content is indexed."""
    from bs4 import BeautifulSoup

    from .discover import content_root

    raw = open(path, "rb").read(3_000_000).decode("utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "header", "footer", "form", "iframe", "svg"]):
        tag.decompose()
    for sel in (".sidebar", ".sidebar-widgets-wrap", "#sidebar", ".breadcrumb", ".breadcrumbs", ".social", ".share", "#google_translate_element"):
        for el in soup.select(sel):
            el.decompose()
    root = soup.select_one("div.entry") if kind == "press_release" else None
    root = root or content_root(soup)
    heading_el = root.find(["h1", "h2"]) or soup.find("h1")
    heading = heading_el.get_text(" ", strip=True) if heading_el else None
    blocks = []
    for el in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "td", "th", "div"], recursive=True):
        if el.name == "div" and el.find(["p", "div", "li", "table", "h1", "h2", "h3", "h4"]):
            continue
        t = el.get_text(" ", strip=True)
        if t and (not blocks or blocks[-1] != t):
            blocks.append(t)
    text = clean("\n".join(blocks))
    return (Extraction(pages=[Page(number=None, text=text)], status="extracted") if text else Extraction(status="empty")), heading
